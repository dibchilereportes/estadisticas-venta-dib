# -*- coding: utf-8 -*-
"""
Informe diario de ventas del CD (DIB + Flooring + DECOEXPRESS) generado desde la plataforma.
Tres vistas (igual que el informe actual): RESUMEN GERENCIAL, B2B-B2C y LINEAS PRODUCTO.

Uso:
    python informe_cd.py AAAA-MM-DD salida.xlsx [--v2026 ruta] [--v2025 ruta] [--datos carpeta]

Fuentes (formato: llave de clasificacion + venta en pesos, agregada por dia):
    --v2026  tabla del ano en curso   (parquet o csv/csv.gz)  [datos/informe_cd_2026.parquet]
    --v2025  tabla del ano anterior   (parquet o csv/csv.gz)  [datos/informe_cd_2025.parquet]
    metas:   datos/metas_cliente_2026.csv y datos/metas_linea_2026.csv
Si la tabla no trae las columnas de clasificacion gerencial, se calculan con
clasificacion_gerencial.py. El script NUNCA envia correos.
"""
import os
import sys
import argparse
import datetime as dt
import numpy as np
import pandas as pd

MES_ES = ['', 'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
          'septiembre', 'octubre', 'noviembre', 'diciembre']
META_MIN = 1_000_000   # si la meta es menor, el % vs meta no es informativo ('s/meta')

# ---------------------------------------------------------------- carga y normalizacion
def _leer(ruta):
    if str(ruta).endswith('.parquet'):
        return pd.read_parquet(ruta)
    return pd.read_csv(ruta)


def _norm(s):
    import unicodedata
    if s is None or (not isinstance(s, str) and pd.isna(s)):
        return ''
    return unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().upper().strip()


def preparar(df, anio_anterior=False):
    """Deja una tabla lista: f, G (grupo), L (linea comercial), bloque, linea, venta, emp."""
    df = df.copy()
    if 'ClasifGerencial' not in df.columns:
        from clasificacion_gerencial import clasificar
        df = df.join(clasificar(df))
    df['f'] = pd.to_datetime(df['x_date'])
    cli = df['x_customer'].map(_norm)
    can = df['x_canal'].map(_norm)
    df['linea'] = df['x_linea'].map(_norm).replace('', 'SIN LINEA')
    es_dexp = df['LineaComercial'].isin(['VENTAS X MAYOR DEXP', 'M PLACE DEXP'])
    df['emp'] = np.where(es_dexp, 'DECOEXPRESS', 'DIB')
    G = df['ClasifGerencial'].copy()
    L = df['LineaComercial'].copy()
    E = df['EntraInforme_CD'].eq('SI')

    vtaemp = pd.Series(False, index=df.index)
    if anio_anterior:
        # 2025 no tenia cuenta analitica que separara B2B/B2C ni venta empresa: se agrupa por
        # cliente (no por la matriz) y el criterio de entrada es "cuenta del CD, no interempresa,
        # no documento no-venta". Venta empresa de flooring en 2025 = facturas de canal OFICINA.
        cuenta = df['x_cuenta_analytica'].map(_norm)
        dexp_all = cuenta.str.contains('DEXP|DECOEXPRES', regex=True)
        en_cd = cuenta.isin(['VXMAYOR B2B', 'VXMAYOR B2C', 'FLOORING B2B', 'FLOORING B2C',
                             'FLOORING VTA EMPRESA']) | dexp_all
        motivo = df['MotivoExclusion'].fillna('')
        E = en_cd & ~(motivo.str.startswith('Interempresa') | motivo.str.startswith('Documento'))
        bra = df['x_branch'].map(_norm)
        def por_nombre(c, k, b):
            txt = c + ' ' + k + ' ' + b
            for key, v in (('CENCOSUD', 'CENCOSUD RETAIL S.A.'), ('PARIS', 'CENCOSUD RETAIL S.A.'),
                           ('FALABELLA', 'FALABELLA RETAIL S.A.'),
                           ('SODIMAC', 'SODIMAC S.A.'), ('EASY', 'EASY RETAIL S.A.'),
                           ('WALMART', 'WALMART CHILE S.A.'), ('WALLMART', 'WALMART CHILE S.A.'),
                           ('ECCSA', 'COM. ECCSA S.A.'), ('RIPLEY', 'COM. ECCSA S.A.'),
                           ('CHILEMAT', 'CHILEMAT SPA.'), ('MATERIALES Y SOLUCIONES', 'MATERIALES Y SOLUCIONES S.A.')):
                if key in txt:
                    return v
            if ('MARKETPLACE' in k or 'MARKETPLACE' in b or 'MERCADO LIBRE' in b
                    or c in ('MERCADO LIBRE', 'HITES S.A.', 'ABC S.A.', 'DUTY FREE TC SA')):
                return 'MARKET PLACE (HITES-LA POLAR-ML)'
            return 'MAYORISTAS'
        G = pd.Series([por_nombre(c_, k_, b_) for c_, k_, b_ in zip(cli, can, bra)], index=df.index)
        G = G.where(~dexp_all, 'DECOEXPRESS')
        flo = cuenta.str.startswith('FLOORING')
        vtaemp = flo & can.eq('OFICINA')
        L = cuenta.where(~vtaemp, 'FLOORING VTA EMPRESA')
        df['emp'] = np.where(dexp_all, 'DECOEXPRESS', 'DIB')
        es_dexp = dexp_all

    # normalizaciones de grupo (decisiones del usuario)
    G = G.replace({'CONSTRUMART S.A.': 'MAYORISTAS'})
    mts = cli.str.contains('MATERIALES Y SOLUCIONES') & L.fillna('').str.startswith('FLOORING')
    G = G.where(~mts, 'MATERIALES Y SOLUCIONES S.A.')
    df['G'] = G
    df['L'] = L
    df['bloque'] = np.where(L.fillna('').str.startswith('FLOORING'), 'FLOORING',
                            np.where(es_dexp, 'DECOEXPRESS', 'DIB'))
    df['E'] = E
    df['vtaemp'] = vtaemp
    return df[df['E']][['f', 'G', 'L', 'bloque', 'linea', 'emp', 'venta', 'vtaemp']]


# ---------------------------------------------------------------- calendario
def habiles(anio, mes, hasta=None):
    import calendar
    import holidays
    fer = holidays.country_holidays('CL', years=[anio])
    n_fer = 0
    total = trans = 0
    for d in range(1, calendar.monthrange(anio, mes)[1] + 1):
        dia = dt.date(anio, mes, d)
        if dia in fer:
            n_fer += 1
        if dia.weekday() < 5 and dia not in fer:
            total += 1
            if hasta and dia <= hasta:
                trans += 1
    return trans, total, n_fer


# ---------------------------------------------------------------- calculo
class Calc:
    def __init__(self, v26, v25, fecha, mc, ml):
        self.v26, self.v25, self.fecha, self.mc, self.ml = v26, v25, fecha, mc, ml
        self.ini = pd.Timestamp(fecha.year, fecha.month, 1)
        self.d = pd.Timestamp(fecha)
        self.ini_ant = pd.Timestamp(fecha.year - 1, fecha.month, 1)
        self.fin_ant = self.ini_ant + pd.offsets.MonthEnd(0)

    def _mask(self, df, **cond):
        m = pd.Series(True, index=df.index)
        for k, v in cond.items():
            m &= df[k].isin(v) if isinstance(v, (list, tuple, set)) else df[k].eq(v)
        return m

    def dia(self, **c):
        df = self.v26
        return df.loc[self._mask(df, **c) & df['f'].eq(self.d), 'venta'].sum()

    def acum(self, **c):
        df = self.v26
        return df.loc[self._mask(df, **c) & df['f'].between(self.ini, self.d), 'venta'].sum()

    def ant(self, excl_ve=True, **c):
        df = self.v25
        m = self._mask(df, **c) & df['f'].between(self.ini_ant, self.fin_ant)
        if excl_ve:                      # venta empresa 2025 queda fuera del ano anterior por cliente
            m &= ~df['vtaemp']
        return df.loc[m, 'venta'].sum()

    def meta_cliente(self, linea=None, grupo=None, cat=None):
        m = self.mc[self.mc['Mes'] == self.fecha.month]
        if linea:
            m = m[m['Linea'] == linea]
        if grupo:
            m = m[m['CategoriaGerencial'] == grupo]
        if cat:
            m = m[m['CategoriaComercial'] == cat]
        return float(m['Meta'].sum())

    def meta_linea(self, emp, linea):
        m = self.ml[(self.ml['Mes'] == self.fecha.month) & (self.ml['Empresa'] == emp) & (self.ml['Linea'] == linea)]
        return float(m['Meta'].sum()) if len(m) else None


def pct(a, base):
    if base is None or base == 0 or pd.isna(base):
        return ''
    r = a / base * 100
    s = '▲' if r >= 100 else ('►' if r >= 90 else '▼')
    return f'{s} {r:.1f}%'.replace('.', ',')


def pct_meta(a, meta):
    if meta is None or meta == 0 or pd.isna(meta):
        return ''
    if meta < META_MIN:
        return 's/meta'
    return pct(a, meta)


# ---------------------------------------------------------------- hojas
def construir(calc):
    R = {'resumen': [], 'b2b2c': [], 'lineas': []}
    c = calc

    def fila(nombre, dia, ac, meta, ant, linea_extra=None):
        return [nombre, dia, ac, meta, ant, pct_meta(ac, meta), pct(ac, ant)]

    # ---------------- RESUMEN
    def grupo_dib(g, nombre, metag=None):
        kw = dict(G=g, bloque='DIB')
        meta = c.meta_cliente('DIB', metag or g)
        return fila(nombre, c.dia(**kw), c.acum(**kw), meta, c.ant(**kw))

    DIB = [('CENCOSUD RETAIL S.A.',) * 2, ('COM. ECCSA S.A.',) * 2, ('EASY RETAIL S.A.',) * 2,
           ('FALABELLA RETAIL S.A.',) * 2, ('SODIMAC S.A.',) * 2, ('WALMART CHILE S.A.',) * 2,
           ('MAYORISTAS', 'MAYORISTAS'), ('MARKET PLACE (HITES-LA POLAR-ML)', 'MARKET PLACE (HITES-LA POLAR-ML)')]
    res = R['resumen']
    res.append(('SEC', 'EMPRESA DIB'))
    dib_rows = [grupo_dib(g, n) for g, n in DIB]
    ch = fila('CHILEMAT SPA.', c.dia(G='CHILEMAT SPA.', bloque='DIB'), c.acum(G='CHILEMAT SPA.', bloque='DIB'), 0,
              c.ant(G='CHILEMAT SPA.', bloque='DIB'))
    if ch[1] or ch[2] or ch[4]:
        dib_rows.append(ch)
    res.extend([['ROW'] + r_ for r_ in dib_rows])

    def total(nombre, rows):
        d = sum(r[1] for r in rows); a = sum(r[2] for r in rows)
        m = sum(r[3] or 0 for r in rows); n = sum(r[4] or 0 for r in rows)
        return fila(nombre, d, a, m, n)
    tot_dib = total('Total DIB', dib_rows)
    res.append(('TOT',) + tuple(tot_dib))

    res.append(('SEC', 'FLOORING'))
    FLO = [('EASY RETAIL S.A.',) * 2, ('SODIMAC S.A.',) * 2, ('MATERIALES Y SOLUCIONES S.A.',) * 2,
           ('CHILEMAT SPA.',) * 2, ('MAYORISTAS', 'MAYORISTAS')]
    flo_rows = []
    for g, n in FLO:
        kw = dict(G=g, bloque='FLOORING')
        flo_rows.append(fila(n, c.dia(**kw), c.acum(**kw), c.meta_cliente('FLOORING', g), c.ant(**kw)))
    res.extend([['ROW'] + r_ for r_ in flo_rows])
    tot_flo = total('Total FLOORING', flo_rows)
    res.append(('TOT',) + tuple(tot_flo))
    res.append(('TOT',) + tuple(total('Total Empresa DIB', dib_rows + flo_rows)))

    res.append(('SEC', 'DECOEXPRESS'))
    kw = dict(bloque='DECOEXPRESS')
    dx = fila('DECOEXPRESS', c.dia(**kw), c.acum(**kw), c.meta_cliente('DECOEXPRESS'), c.ant(**kw))
    res.append(tuple(['ROW'] + dx))
    res.append(('TOT',) + tuple(['Total DECOEXPRESS'] + dx[1:]))

    # ---------------- B2B-B2C
    b = R['b2b2c']
    def bloque_fila(titulo, lineas_c, clientes, cat_prefix, linea_meta):
        rows = []
        for cli_g, etiqueta, cat in clientes:
            kw = dict(G=cli_g, L=lineas_c)
            meta = c.meta_cliente(linea_meta, cli_g, cat) if cat else None
            rows.append(fila(etiqueta, c.dia(**kw), c.acum(**kw), meta, None) + [titulo])
        return rows
    def pbloque(titulo, filas):
        b.append(('SEC', titulo))
        for f in filas:
            b.append(('ROW', f[0], f[7] if len(f) > 7 else '', f[1], f[2], f[3], '', f[5], ''))
    def emitir(titulo_total, items):
        # items: lista de (etiqueta, linea_comercial, grupo, linea_meta, categoria)
        subtotal = [0, 0, 0]
        for et, lc, g, lm, cat in items:
            d = c.dia(G=g, L=lc); a = c.acum(G=g, L=lc)
            meta = c.meta_cliente(lm, g, cat) if cat else None
            if d == 0 and a == 0 and not meta:
                continue
            b.append(('ROW', et, lc, d, a, meta, '', pct_meta(a, meta), ''))
            subtotal[0] += d; subtotal[1] += a; subtotal[2] += meta or 0
        b.append(('TOT', titulo_total, '', subtotal[0], subtotal[1], subtotal[2] or None, '', pct_meta(subtotal[1], subtotal[2]), ''))
        return subtotal
    tb = [0, 0, 0]; tc = [0, 0, 0]
    b.append(('SEC', 'VENTAS B2B'))
    s = emitir('TOTAL VXMAYOR B2B', [
        ('EASY RETAIL S.A.', 'VXMAYOR B2B', 'EASY RETAIL S.A.', 'DIB', 'B2B EASY'),
        ('SODIMAC S.A.', 'VXMAYOR B2B', 'SODIMAC S.A.', 'DIB', 'B2B SODIMAC'),
        ('MAYORISTAS', 'VXMAYOR B2B', 'MAYORISTAS', 'DIB', 'B2B MAYORISTAS'),
        ('CHILEMAT SPA.', 'VXMAYOR B2B', 'CHILEMAT SPA.', 'DIB', None)])
    tb = [x + y for x, y in zip(tb, s)]
    s = emitir('TOTAL FLOORING B2B', [
        ('EASY RETAIL S.A.', 'FLOORING B2B', 'EASY RETAIL S.A.', 'FLOORING', None),
        ('SODIMAC S.A.', 'FLOORING B2B', 'SODIMAC S.A.', 'FLOORING', None),
        ('MATERIALES Y SOLUCIONES S.A.', 'FLOORING B2B', 'MATERIALES Y SOLUCIONES S.A.', 'FLOORING', 'B2B MTS'),
        ('CHILEMAT SPA.', 'FLOORING B2B', 'CHILEMAT SPA.', 'FLOORING', 'B2B CHILEMAT'),
        ('MAYORISTAS', 'FLOORING B2B', 'MAYORISTAS', 'FLOORING', 'B2B MAYORISTAS')])
    tb = [x + y for x, y in zip(tb, s)]
    s = emitir('TOTAL FLOORING VTA EMPRESA', [
        ('SODIMAC S.A.', 'FLOORING VTA EMPRESA', 'SODIMAC S.A.', 'FLOORING', None),
        ('EASY RETAIL S.A.', 'FLOORING VTA EMPRESA', 'EASY RETAIL S.A.', 'FLOORING', None),
        ('MAYORISTAS', 'FLOORING VTA EMPRESA', 'MAYORISTAS', 'FLOORING', None)])
    tb = [x + y for x, y in zip(tb, s)]
    b.append(('TOT', 'TOTAL B2B', '', tb[0], tb[1], tb[2] or None, '', pct_meta(tb[1], tb[2]), ''))
    b.append(('SEC', 'VENTAS B2C'))
    s = emitir('TOTAL VXMAYOR B2C', [
        ('CENCOSUD RETAIL S.A.', 'VXMAYOR B2C', 'CENCOSUD RETAIL S.A.', 'DIB', 'B2C CENCOSUD'),
        ('COM. ECCSA S.A.', 'VXMAYOR B2C', 'COM. ECCSA S.A.', 'DIB', 'B2C RIPLEY'),
        ('EASY RETAIL S.A.', 'VXMAYOR B2C', 'EASY RETAIL S.A.', 'DIB', 'B2C EASY'),
        ('FALABELLA RETAIL S.A.', 'VXMAYOR B2C', 'FALABELLA RETAIL S.A.', 'DIB', 'B2C FALABELLA'),
        ('SODIMAC S.A.', 'VXMAYOR B2C', 'SODIMAC S.A.', 'DIB', 'B2C SODIMAC'),
        ('WALMART CHILE S.A.', 'VXMAYOR B2C', 'WALMART CHILE S.A.', 'DIB', 'B2C WALMART'),
        ('MARKET PLACE (HITES-LA POLAR-ML)', 'VXMAYOR B2C', 'MARKET PLACE (HITES-LA POLAR-ML)', 'DIB', 'B2C ML')])
    tc = [x + y for x, y in zip(tc, s)]
    s = emitir('TOTAL FLOORING B2C', [
        ('EASY RETAIL S.A.', 'FLOORING B2C', 'EASY RETAIL S.A.', 'FLOORING', 'B2C EASY'),
        ('SODIMAC S.A.', 'FLOORING B2C', 'SODIMAC S.A.', 'FLOORING', 'B2C SODIMAC')])
    tc = [x + y for x, y in zip(tc, s)]
    b.append(('TOT', 'TOTAL B2C', '', tc[0], tc[1], tc[2] or None, '', pct_meta(tc[1], tc[2]), ''))
    b.append(('TOT', 'TOTAL EMPRESA DIB', '', tb[0] + tc[0], tb[1] + tc[1], (tb[2] + tc[2]) or None, '',
              pct_meta(tb[1] + tc[1], tb[2] + tc[2]), ''))
    b.append(('SEC', 'DECOEXPRESS'))
    for et, lc, cat in (('DECOEXPRESS B2B', 'VENTAS X MAYOR DEXP', 'DECOEXPRESS B2B'),
                        ('DECOEXPRESS B2C', 'M PLACE DEXP', 'DECOEXPRESS B2C')):
        d = c.dia(L=lc); a = c.acum(L=lc); m = c.meta_cliente('DECOEXPRESS', None, cat)
        b.append(('ROW', et, et, d, a, m, '', pct_meta(a, m), ''))
    d = c.dia(emp='DECOEXPRESS'); a = c.acum(emp='DECOEXPRESS'); m = c.meta_cliente('DECOEXPRESS')
    b.append(('TOT', 'TOTAL DECOEXPRESS', '', d, a, m, '', pct_meta(a, m), ''))

    # ---------------- LINEAS PRODUCTO
    ln = R['lineas']
    def seccion(emp, titulo, orden):
        ln.append(('SEC', titulo))
        visto = set(); tot = [0, 0, 0, 0]
        v = c.v26[c.v26['emp'] == emp]
        extra = [x for x in sorted(set(v['linea'])) if x not in orden and v.loc[v['linea'] == x, 'venta'].abs().sum() > 0]
        for l in orden + extra:
            d = c.dia(emp=emp, linea=l); a = c.acum(emp=emp, linea=l); n = c.ant(excl_ve=False, emp=emp, linea=l)
            m = c.meta_linea(emp, l)
            if not (d or a or n or m):
                continue
            ln.append(('ROW', l, d, a, m, n, pct_meta(a, m), pct(a, n)))
            tot[0] += d; tot[1] += a; tot[2] += m or 0; tot[3] += n
        ln.append(('TOT', 'TOTAL ' + titulo.replace('EMPRESA ', ''), tot[0], tot[1], tot[2] or None, tot[3],
                   pct_meta(tot[1], tot[2]), pct(tot[1], tot[3])))
    seccion('DIB', 'EMPRESA DIB', ['ALFOMBRAS', 'LIMPIAPIES', 'TEXTIL HOGAR', 'FLOORING', 'MUEBLES', 'SERVICIOS', 'HOTELERIA'])
    seccion('DECOEXPRESS', 'DECOEXPRESS', ['MENAJE', 'BANO', 'DECORACION', 'MUEBLES', 'SERVICIOS', 'SIN LINEA', 'COCINA'])
    return R


# ---------------------------------------------------------------- Excel
def escribir(R, fecha, salida):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    AZUL, CLARO, GRIS = '1F3A5F', 'DCE6F2', 'F2F2F2'
    hf = Font(bold=True, color='FFFFFF'); hfill = PatternFill('solid', fgColor=AZUL)
    sfill = PatternFill('solid', fgColor=CLARO); tfill = PatternFill('solid', fgColor=GRIS)
    fino = Side(style='thin', color='BFBFBF')
    NUM = '#,##0;[Red]-#,##0'
    wb = Workbook()
    trans, total, n_fer = habiles(fecha.year, fecha.month, fecha)

    def cabecera(ws, fila, cols):
        for j, t in enumerate(cols, 1):
            c = ws.cell(fila, j, t); c.font = hf; c.fill = hfill
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def pintar(ws, r, vals, estilo, ncols_num):
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.border = Border(bottom=fino)
            if estilo == 'SEC':
                c.fill = sfill; c.font = Font(bold=True)
            elif estilo == 'TOT':
                c.fill = tfill; c.font = Font(bold=True)
            if j in ncols_num and isinstance(v, (int, float, np.floating, np.integer)):
                c.number_format = NUM
            if j > 1 and isinstance(v, str) and v[:1] in '▲►▼s':
                c.alignment = Alignment(horizontal='center')
                if v[:1] == '▲':
                    c.font = Font(bold=estilo == 'TOT', color='1B7F3B')
                elif v[:1] == '▼':
                    c.font = Font(bold=estilo == 'TOT', color='B3261E')
                elif v[:1] == '►':
                    c.font = Font(bold=estilo == 'TOT', color='9A6B00')

    # --- hoja 1
    ws = wb.active; ws.title = 'RESUMEN GERENCIAL'
    ws['A1'] = 'Informe de ventas CD'; ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = 'Fecha del informe'; ws['B2'] = fecha.strftime('%d-%m-%Y')
    ws['A3'] = 'Día hábil transcurrido'; ws['B3'] = trans
    ws['A4'] = 'Días hábiles del mes'; ws['B4'] = total
    ws['A5'] = 'Feriados del mes'; ws['B5'] = n_fer
    cab = ['Cliente / Segmento', 'Venta Día', 'Venta Acumulada', 'Meta Mensual', 'Igual Mes Año Ant.', '% vs Meta', '% vs Año Ant.']
    r = 7
    cabecera(ws, r, cab); r += 1
    for item in R['resumen']:
        tipo = item[0]
        if tipo == 'SEC':
            pintar(ws, r, [item[1]] + [''] * 6, 'SEC', ());
        else:
            pintar(ws, r, list(item[1:]), 'TOT' if tipo == 'TOT' else 'ROW', (2, 3, 4, 5))
        r += 1
    ws.column_dimensions['A'].width = 38
    for col in 'BCDEFG':
        ws.column_dimensions[col].width = 18

    # --- hoja 2
    ws2 = wb.create_sheet('B2B-B2C')
    cab2 = ['Clasificación', 'Línea Comercial', 'Venta Día', 'Venta Acumulada', 'Meta Mensual', '', '% vs Meta']
    cabecera(ws2, 1, cab2); r = 2
    for item in R['b2b2c']:
        tipo = item[0]
        if tipo == 'SEC':
            pintar(ws2, r, [item[1]] + [''] * 6, 'SEC', ())
        else:
            v = list(item[1:])           # etiqueta, linea, dia, acum, meta, '', pct, ''
            pintar(ws2, r, [v[0], v[1], v[2], v[3], v[4], '', v[6]], 'TOT' if tipo == 'TOT' else 'ROW', (3, 4, 5))
        r += 1
    ws2.cell(r + 1, 1, 'Nota: "Igual mes año anterior" no se muestra en esta vista: en 2025 la cuenta analítica no separaba B2B/B2C; se informa por cliente y por línea en las otras hojas.').font = Font(italic=True, color='666666')
    ws2.column_dimensions['A'].width = 38; ws2.column_dimensions['B'].width = 24
    for col in 'CDEFG':
        ws2.column_dimensions[col].width = 18

    # --- hoja 3
    ws3 = wb.create_sheet('LINEAS PRODUCTO')
    cab3 = ['Línea de Producto', 'Venta Día', 'Venta Acumulada', 'Meta Mensual', 'Igual Mes Año Ant.', '% vs Meta', '% vs Año Ant.']
    cabecera(ws3, 1, cab3); r = 2
    for item in R['lineas']:
        tipo = item[0]
        if tipo == 'SEC':
            pintar(ws3, r, [item[1]] + [''] * 6, 'SEC', ())
        else:
            pintar(ws3, r, list(item[1:]), 'TOT' if tipo == 'TOT' else 'ROW', (2, 3, 4, 5))
        r += 1
    ws3.column_dimensions['A'].width = 28
    for col in 'BCDEFG':
        ws3.column_dimensions[col].width = 18
    for w in (ws, ws2, ws3):
        w.sheet_view.showGridLines = False
    ws.freeze_panes = 'B8'; ws2.freeze_panes = 'A2'; ws3.freeze_panes = 'A2'
    wb.save(salida)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('fecha'); ap.add_argument('salida')
    ap.add_argument('--datos', default='datos')
    ap.add_argument('--v2026'); ap.add_argument('--v2025')
    a = ap.parse_args()
    fecha = dt.date.fromisoformat(a.fecha)
    v26 = preparar(_leer(a.v2026 or os.path.join(a.datos, 'informe_cd_2026.parquet')))
    v25 = preparar(_leer(a.v2025 or os.path.join(a.datos, 'informe_cd_2025.parquet')), anio_anterior=True)
    mc = pd.read_csv(os.path.join(a.datos, 'metas_cliente_2026.csv'))
    ml = pd.read_csv(os.path.join(a.datos, 'metas_linea_2026.csv'))
    R = construir(Calc(v26, v25, fecha, mc, ml))
    escribir(R, fecha, a.salida)
    print('Informe generado:', a.salida)


if __name__ == '__main__':
    main()
